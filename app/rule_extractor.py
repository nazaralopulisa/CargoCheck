"""
rule_extractor.py - reads the 7 fields from SI and BL documents with plain code,
no AI. Returns exactly the same shape as extractor.py (the Gemini/Bedrock one),
so the pipeline can try this first and only call the LLM when it's not enough.

How it works:
  1. Get the document's text (.txt directly; PDF/Word/Excel via a library).
  2. Work out what the document is from its title (SI, BL, or something else).
  3. Find every KNOWN label ("Load Port:", "POL:", "Freight:" ...). A value is
     the text between its label and the next known label.
  4. Map labels to our 7 fields. Labels we don't compare (Freight, Vessel...)
     are still listed, so they correctly END the value before them.

Traps handled on purpose (found by surveying the real attachments):
  - "NET WEIGHT" is never read as gross weight
  - "To the Order of:", "P.O. BOX:", "TEL:", "PIN NO.:" live INSIDE addresses,
    so they are NOT labels and don't cut a consignee value short
  - longest label wins: "Notify Party/Intermediate Consignee" is notify_party,
    not consignee
  - labels with extra non-English text still count: "Gross Weight毛重(KGS):"
  - company names that continue on the next line are joined consistently:
    "APRIL FINE PAPER TRADING" + "ON BEHALF OF VITAL SOLUTIONS PTE LTD; 77 ROBINSON..."
    -> "APRIL FINE PAPER TRADING ON BEHALF OF VITAL SOLUTIONS PTE LTD"
"""

import io
import re
from pathlib import Path

FIELDS = ["shipper", "consignee", "notify_party", "port_of_loading",
          "port_of_discharge", "container_count", "gross_weight_kg"]

# --- Step 1: the label dictionary ---------------------------------------------
# Each entry is a regex. Matching ignores case and spacing differences.

FIELD_LABELS = {
    "shipper": [r"shipper\s*\(principal\s*or\s*seller\)", r"shipper\s*/\s*exporter",
                r"shipper", r"exporter"],
    "consignee": [r"consignee\s*\(non-negotiable\)", r"consignee"],
    "notify_party": [r"notify\s*party\s*/\s*intermediate\s*consignee",
                     r"notify\s*party", r"notify"],
    "port_of_loading": [r"port\s*of\s*loading\s*\(pol\)", r"port\s*of\s*loading",
                        r"load(?:ing)?\s*port", r"pol"],
    "port_of_discharge": [r"port\s*of\s*discharge\s*\(pod\)", r"port\s*of\s*discharge",
                          r"discharge\s*port", r"pod"],
    "container_count": [r"no\.?\s*of\s*containers\s*or\s*packages", r"no\.?\s*of\s*containers",
                        r"total\s*containers", r"container\s*count"],
    "gross_weight_kg": [r"gross\s*(?:weight|wt)\.?(?:\s*\([^)]*\))?"],
}

# Labels we don't compare, but which must END the value before them.
OTHER_LABELS = [
    r"net\s*(?:weight|wt)\.?(?:\s*\([^)]*\))?", r"freight", r"description\s*of\s*goods",
    r"description", r"commodity", r"oc\s*no\.?", r"hs\s*code",
    r"booking\s*(?:reference|ref\.?|no\.?)", r"ocean\s*vessel", r"vessel\s*name", r"vessel",
    r"voyage\s*no\.?", r"voyage", r"voy\.?\s*no\.?", r"voy\.?",
    r"bill\s*of\s*lading\s*no\.?", r"b\s*/?\s*l\s*(?:no\.?|number)", r"country\s*of\s*origin",
    r"certificate\s*no\.?", r"issuing\s*authority", r"invoice\s*(?:no\.?|date)",
    r"seller", r"buyer", r"total\s*amount", r"payment\s*terms", r"incoterms",
    r"place\s*of\s*(?:receipt|delivery)", r"measurement", r"marks\s*(?:and|&)\s*numbers",
    r"date", r"new\s*no\.?",
]

_ALL = [(f, p) for f, pats in FIELD_LABELS.items() for p in pats] + [(None, p) for p in OTHER_LABELS]
_ALL.sort(key=lambda fp: len(fp[1]), reverse=True)          # longest label first
# After the label we allow extra non-English text and a unit in brackets before
# the colon, e.g. "Gross Weight毛重(KGS):". Only the English label is captured.
LABEL_RE = re.compile(r"(?<![A-Za-z])(" + "|".join(p for _, p in _ALL) + r")"
                      r"(?:\s*[^\x00-\x7F]+)?(?:\s*\([^)]*\))?\s*:", re.I)
_COMPILED = [(f, re.compile(p + r"$", re.I)) for f, p in _ALL]

# Document titles. "Bill of Lading No." is a LABEL, not a title, so it's excluded.
TITLE_PATTERNS = [
    ("SI", re.compile(r"SHIPPING\s+INSTRUCTIONS?", re.I)),
    ("BL", re.compile(r"BILL\s+OF\s+LADING(?!\s*(?:NO\b|NO\.|NUMBER|#))", re.I)),
    ("BL", re.compile(r"SEA\s*WAYBILL", re.I)),
    # "Other" titles only count when they stand on a line of their own. A real SI
    # often mentions "3 Original invoice, 3 Packing list" in a field, and a BL may
    # have "Invoice No.:" - those words must not turn it into OTHER.
    ("OTHER", re.compile(r"^[\s=*#-]*(?:COMMERCIAL\s+INVOICE|PROFORMA\s+INVOICE|INVOICE|"
                         r"CERTIFICATE\s+OF\s+ORIGIN|PACKING\s+LIST|CERTIFICATE)"
                         r"[\s()A-Z]{0,20}$", re.I | re.M)),
]


def which_field(label_text):
    label_text = label_text.strip()
    for field, rx in _COMPILED:
        if rx.match(label_text):
            return field
    return None


# --- Step 2: get text out of any file type -----------------------------------

def rows_to_text(rows):
    """Table rows -> 'label: value' lines, so the same label rules work."""
    lines = []
    for cells in rows:
        cells = [str(c).strip() for c in cells if c is not None and str(c).strip()]
        if len(cells) >= 2 and not cells[0].endswith(":"):
            lines.append(f"{cells[0]}: {'  '.join(cells[1:])}")
        elif cells:
            lines.append("  ".join(cells))
    return "\n".join(lines)


def document_text(inbox, path):
    """Return (text, problem). problem is None when the text was read fine."""
    ext = Path(path).suffix.lower()
    try:
        if ext == ".txt":
            return inbox.read_text(path), None
        data = inbox.read_bytes(path)
        if ext == ".pdf":
            import pdfplumber
            with pdfplumber.open(io.BytesIO(data)) as pdf:
                text = "\n".join(page.extract_text(use_text_flow=True) or "" for page in pdf.pages)
            if not text.strip():
                return "", "scanned or image-only PDF, no text layer"
            return text, None
        if ext == ".docx":
            import docx
            d = docx.Document(io.BytesIO(data))
            parts = [p.text for s in d.sections for p in s.header.paragraphs]  # titles often live here
            parts += [p.text for p in d.paragraphs]
            for table in d.tables:
                parts.append(rows_to_text([c.text for c in row.cells] for row in table.rows))
            return "\n".join(parts), None
        if ext == ".xlsx":
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
            # the sheet name often holds the title, e.g. a tab called "SHIPPING INSTRUCTION"
            return "\n".join(f"{ws.title}\n" + rows_to_text(ws.iter_rows(values_only=True))
                             for ws in wb.worksheets), None
        return "", f"unsupported file type {ext}"
    except Exception as e:
        return "", f"could not open: {e}"


# --- Step 3: document type + field values ------------------------------------

def detect_type(text, path):
    """Returns (type, sure). The title that appears FIRST in the document wins.
    sure=False means no title was found and we guessed from the filename."""
    head = text[:1500]
    found = [(m.start(), kind) for kind, rx in TITLE_PATTERNS for m in [rx.search(head)] if m]
    if found:
        return min(found)[1], True
    name = Path(path).stem.upper()
    if name.endswith("_SI"):
        return "SI", False
    if name.endswith("_BL"):
        return "BL", False
    return "OTHER", False


def filename_type(path):
    name = Path(path).stem.upper()
    return "SI" if name.endswith("_SI") else "BL" if name.endswith("_BL") else None


def first_piece(value):
    """First real line of a value; also splits text that was flattened onto one
    line (wide gaps of 2+ spaces separate the name from the address)."""
    for line in value.splitlines():
        line = line.strip(" \t;,")
        if line:
            return re.split(r"\s{2,}", line)[0].strip(" ;,")
    return None


PARTY_FIELDS = {"shipper", "consignee", "notify_party"}
NAME_CONTINUATION = re.compile(r"^(ON BEHALF OF|FOR AND ON BEHALF|C/O|A/C|\()", re.I)
COMPANY_SUFFIX = re.compile(
    r"\b(LTD|LIMITED|PTE|SDN BHD|BHD|FZE|FZCO|FZ-LLC|LLC|INC|CORP|GMBH|CO)\b", re.I)


def party_name(value):
    """Company name, including a second line when it continues the name.

    Values look like:  NAME            (first line)
                         CONTINUATION; ADDRESS; ADDRESS   (indented next line)
    or the same thing flattened onto one line with wide gaps between the parts.
    The continuation is kept only if it reads like part of a name
    ("ON BEHALF OF ...", "(MIDDLE EAST) FZE"), never an address.
    """
    pieces = []
    for line in value.splitlines():
        for chunk in re.split(r"\s{2,}", line.strip()):
            if chunk.strip(" ;,"):
                pieces.append(chunk.strip())
    if not pieces:
        return None

    name = pieces[0].split(";")[0].strip(" ;,")
    if ";" in pieces[0] or len(pieces) < 2:
        return name or None                 # name ended at ";" or there is nothing after it

    nxt = pieces[1].split(";")[0].strip(" ;,")
    looks_like_name = NAME_CONTINUATION.match(nxt) or (
        COMPANY_SUFFIX.search(nxt) and not re.search(r"\d", nxt))
    return f"{name} {nxt}" if looks_like_name else name


def extract_fields(text):
    matches = list(LABEL_RE.finditer(text))
    fields = {f: None for f in FIELDS}
    for i, m in enumerate(matches):
        field = which_field(m.group(1))
        if field is None or fields[field] is not None:       # first occurrence wins
            continue
        end = matches[i + 1].start() if i + 1 < len(matches) else len(text)
        raw = text[m.end():end]
        fields[field] = party_name(raw) if field in PARTY_FIELDS else first_piece(raw)
    return fields


def extract_document(inbox, path):
    text, problem = document_text(inbox, path)
    if problem:
        return None, problem
    doc_type, sure = detect_type(text, path)
    return {"doc_type_detected": doc_type, "fields": extract_fields(text),
            "source_file": path, "method": "rules", "type_from_title": sure}, None


# --- Step 4: one email, same output shape as extractor.py --------------------

def extract_email(inbox, email):
    result = {"email_id": email["email_id"], "si": None, "bl": None,
              "other_documents": [], "extraction_issues": [], "method": "rules"}
    attachments = email.get("attachments") or []
    if not attachments:
        result["extraction_issues"].append("No attachments on this email")
        return result

    docs = []
    for path in attachments:
        doc, problem = extract_document(inbox, path)
        if problem:
            result["extraction_issues"].append(f"{path} is empty or unreadable ({problem})")
        else:
            docs.append(doc)

    # Two documents claim the same slot (e.g. an SI that mentions "bill of lading"
    # near the top) while the other slot is empty: trust the filename to break the tie.
    for kind, other in (("SI", "BL"), ("BL", "SI")):
        same = [d for d in docs if d["doc_type_detected"] == kind]
        if len(same) > 1 and not any(d["doc_type_detected"] == other for d in docs):
            for d in same:
                if filename_type(d["source_file"]) == other:
                    d["doc_type_detected"], d["type_from_title"] = other, False
                    break

    for doc in docs:
        kind, path = doc["doc_type_detected"], doc["source_file"]
        if kind == "SI" and result["si"] is None:
            result["si"] = doc
        elif kind == "BL" and result["bl"] is None:
            result["bl"] = doc
        else:
            result["other_documents"].append(doc)
            result["extraction_issues"].append(f"{path} was detected as {kind}, not a usable SI/BL")

    if result["si"] is None:
        result["extraction_issues"].append("No SI found")
    if result["bl"] is None:
        result["extraction_issues"].append("No BL found")
    return result


def is_confident(result):
    """True when the rules did the job and an LLM wouldn't do better:
    every document was identified from its title, and every SI/BL found has
    all 7 fields. Otherwise the pipeline should ask the LLM."""
    docs = [d for d in (result.get("si"), result.get("bl")) if d] + result.get("other_documents", [])
    if not docs or any("unreadable" in i for i in result.get("extraction_issues", [])):
        return False
    # Missing SI or BL is exactly when a second opinion matters most: let the LLM check
    # whether a document the rules called OTHER is really an SI or BL.
    if result.get("si") is None or result.get("bl") is None:
        return False
    if not all(d.get("type_from_title") for d in docs):
        return False
    return all(all(v is not None for v in d["fields"].values())
               for d in (result.get("si"), result.get("bl")) if d)


# --- Step 5: self-test ---------------------------------------------------------

if __name__ == "__main__":
    class FakeInbox:
        def __init__(self, files):
            self.files = files
        def read_text(self, p):
            return self.files[p]
        def read_bytes(self, p):
            v = self.files[p]
            return v if isinstance(v, bytes) else v.encode()

    si_flat = ("SHIPPING INSTRUCTION ========================================  Shipper (Principal or Seller): "
               "ASIA PACIFIC PAPERBOARD TRADING PTE LTD   80 RAFFLES PLACE, #50-01 UOB PLAZA 1; SINGAPORE 048624 "
               "CONSIGNEE: PACIFIC OFFICE (M) SDN BHD   LOT 6, JALAN P/7; SECTION 13, 43650 BANDAR BARU BANGI; "
               "SELANGOR, MALAYSIA NOTIFY PARTY: PACIFIC OFFICE (M) SDN BHD Load Port: PORT KLANG (WESTPORT), "
               "MALAYSIA (MYPKG) Discharge Port: MERSIN, TURKEY (TRMER) No. of Containers or Packages: 1 x 40'HC "
               "Gross Wt (kgs): 20,381 KG Vessel: NAP 914 V.BS007 Voy.: QI540A Description: COATED IVORY BOARD "
               "HS Code: 48105900 Booking Reference: I978820812 OC No.: 5ALT-19136 Freight: PREPAID")
    bl = """BILL OF LADING (DRAFT)
========================================
Shipper/Exporter: ASIA PACIFIC PAPERBOARD TRADING PTE LTD
  80 RAFFLES PLACE, #50-01 UOB PLAZA 1; SINGAPORE 048624
Consignee (Non-Negotiable): To the Order of: CIMB BANK BERHAD
  P.O. BOX: 10338, KUALA LUMPUR
Notify Party/Intermediate Consignee: PACIFIC OFFICE (M) SDN BHD
POL: PORT KLANG (WESTPORT), MALAYSIA (MYPKG)
Port of Discharge (POD): MERSIN, TURKEY (TRMER)
Total Containers: 1
NET WEIGHT: 19,900 KG
GROSS WEIGHT (KGS): 20,381
Freight: PREPAID
"""
    invoice = "COMMERCIAL INVOICE\nInvoice No.: 123\nSeller: X\nBuyer: Y\nTotal Amount: 5000"

    # a Word file with a two-column table, and an Excel sheet
    import docx, openpyxl
    d = docx.Document(); d.add_paragraph("BILL OF LADING")
    t = d.add_table(rows=0, cols=2)
    for k, v in [("Shipper", "ASIA PACIFIC PAPERBOARD TRADING PTE LTD"), ("Consignee", "PACIFIC OFFICE (M) SDN BHD"),
                 ("Notify", "PACIFIC OFFICE (M) SDN BHD"), ("Load Port", "PORT KLANG (MYPKG)"),
                 ("POD", "MERSIN (TRMER)"), ("Container Count", "4"), ("Gross Weight (KG)", "20381")]:
        c = t.add_row().cells; c[0].text, c[1].text = k, v
    buf = io.BytesIO(); d.save(buf)
    wb = openpyxl.Workbook(); ws = wb.active; ws.append(["SHIPPING INSTRUCTION"])
    for k, v in [("SHIPPER", "ASIA PACIFIC PAPERBOARD TRADING PTE LTD"), ("CONSIGNEE", "PACIFIC OFFICE (M) SDN BHD"),
                 ("NOTIFY PARTY", "PACIFIC OFFICE (M) SDN BHD"), ("PORT OF LOADING", "PORT KLANG (MYPKG)"),
                 ("PORT OF DISCHARGE", "MERSIN (TRMER)"), ("No. of Containers", "3 x 40HC"), ("GROSS WEIGHT", "20,381 KGS")]:
        ws.append([k, v])
    xbuf = io.BytesIO(); wb.save(xbuf)

    inbox = FakeInbox({"a/e1_SI.txt": si_flat, "a/e1_BL.txt": bl, "a/e2_BL.txt": invoice,
                       "a/e3_BL.docx": buf.getvalue(), "a/e3_SI.xlsx": xbuf.getvalue()})

    checks = []
    r1 = extract_email(inbox, {"email_id": "e1", "attachments": ["a/e1_SI.txt", "a/e1_BL.txt"]})
    si, blf = r1["si"]["fields"], r1["bl"]["fields"]
    checks += [
        ("flattened SI: shipper name only", si["shipper"] == "ASIA PACIFIC PAPERBOARD TRADING PTE LTD"),
        ("flattened SI: consignee name only", si["consignee"] == "PACIFIC OFFICE (M) SDN BHD"),
        ("flattened SI: notify stops at next label", si["notify_party"] == "PACIFIC OFFICE (M) SDN BHD"),
        ("flattened SI: port with code", si["port_of_loading"] == "PORT KLANG (WESTPORT), MALAYSIA (MYPKG)"),
        ("flattened SI: containers", si["container_count"] == "1 x 40'HC"),
        ("flattened SI: weight", si["gross_weight_kg"] == "20,381 KG"),
        ("'To the Order of' kept inside consignee", blf["consignee"] == "To the Order of: CIMB BANK BERHAD"),
        ("longest label wins for notify party", blf["notify_party"] == "PACIFIC OFFICE (M) SDN BHD"),
        ("NET WEIGHT ignored, gross used", blf["gross_weight_kg"] == "20,381"),
        ("SI+BL fully read -> confident", is_confident(r1)),
    ]
    r2 = extract_email(inbox, {"email_id": "e2", "attachments": ["a/e1_SI.txt", "a/e2_BL.txt"]})
    checks += [("invoice named _BL detected as OTHER", r2["bl"] is None and
                r2["other_documents"][0]["doc_type_detected"] == "OTHER")]
    r3 = extract_email(inbox, {"email_id": "e3", "attachments": ["a/e3_SI.xlsx", "a/e3_BL.docx"]})
    checks += [("Excel SI read", r3["si"] and r3["si"]["fields"]["container_count"] == "3 x 40HC"),
               ("Word BL table read", r3["bl"] and r3["bl"]["fields"]["container_count"] == "4"),
               ("Word/Excel fully read -> confident", is_confident(r3))]
    si_with_bl_label = "Ref / Bill of Lading No.: TBA\n" + si_flat
    inbox.files["a/e5_SI.txt"] = si_with_bl_label
    r5 = extract_email(inbox, {"email_id": "e5", "attachments": ["a/e5_SI.txt", "a/e1_BL.txt"]})
    checks += [("'Bill of Lading No.' label doesn't make an SI a BL",
                r5["si"] is not None and r5["bl"] is not None and not r5["other_documents"])]
    swapped = "BILL OF LADING\n" + si_flat.replace("SHIPPING INSTRUCTION", "")
    inbox.files["a/e6_SI.txt"] = swapped
    r6 = extract_email(inbox, {"email_id": "e6", "attachments": ["a/e6_SI.txt", "a/e1_BL.txt"]})
    checks += [("two BLs, one named _SI -> filename fills the SI slot",
                r6["si"] is not None and r6["bl"] is not None)]
    r4 = extract_email(inbox, {"email_id": "e4", "attachments": []})
    checks += [("no attachments -> not confident", not is_confident(r4))]

    # real layouts found in the data (email_044-style SI/BL, flattened text, Chinese label)
    bl044 = """BILL OF LADING (DRAFT)
Shipper/Exporter: APRIL FINE PAPER TRADING
  ON BEHALF OF VITAL SOLUTIONS PTE LTD; 77 ROBINSON ROAD, #21-01; SINGAPORE 068896
CONSIGNEE: SAFQA LIMITED
  P.O. BOX 99423-80100; TONONOKA ROAD; MOMBASA, KENYA; PIN NO.: P051376597X
Notify Party: SAFQA LIMITED
Gross Weight毛重(KGS): 47,192 KG
Ocean Vessel: INDO SUKSES 65 V.51NW1
"""
    f044 = extract_fields(bl044)
    flat044 = extract_fields(" ".join(bl044.split("\n")).replace("TRADING ", "TRADING   ", 1))
    me = extract_fields("Shipper: APRIL FINE PAPER TRADING\n  (MIDDLE EAST) FZE; #813, 4 EA, DUBAI\nFreight: X")
    kr = extract_fields("Consignee: MOORIM SP CO., LTD\n  656, GANGNAM-DAERO, GANGNAM-GU; SEOUL\nFreight: X")
    checks += [
        ("name continued on next line is joined",
         f044["shipper"] == "APRIL FINE PAPER TRADING ON BEHALF OF VITAL SOLUTIONS PTE LTD"),
        ("same name when flattened onto one line",
         flat044["shipper"] == "APRIL FINE PAPER TRADING ON BEHALF OF VITAL SOLUTIONS PTE LTD"),
        ("'(MIDDLE EAST) FZE' continuation joined",
         me["shipper"] == "APRIL FINE PAPER TRADING (MIDDLE EAST) FZE"),
        ("P.O. BOX address not joined", f044["consignee"] == "SAFQA LIMITED"),
        ("street address not joined", kr["consignee"] == "MOORIM SP CO., LTD"),
        ("label with Chinese text still read", f044["gross_weight_kg"] == "47,192 KG"),
    ]

    # Excel/Word layouts that were wrongly called OTHER
    def xlsx(title, rows):
        wb2 = openpyxl.Workbook(); ws2 = wb2.active; ws2.title = title
        for r in rows:
            ws2.append(r)
        b = io.BytesIO(); wb2.save(b); return b.getvalue()
    xrows = [["Shipper", "APRIL FAR EAST (M) SDN BHD"], ["Consignee", "MOORIM SP CO., LTD"],
             ["Notify Party", "MOORIM SP CO., LTD"], ["Port of Loading", "PORT KLANG (MYPKG)"],
             ["Port of Discharge", "CALLAO (PECLL)"], ["No. of Containers", "3 x 40'HC"],
             ["Gross Weight", "22,000 KG"], ["Documents Required", "3 Original invoice, 3 Packing list"]]
    d2 = docx.Document(); d2.sections[0].header.paragraphs[0].text = "BILL OF LADING"
    t2 = d2.add_table(rows=0, cols=2)
    for k, v in [["Invoice No.", "INV-1"]] + xrows:
        c = t2.add_row().cells; c[0].text, c[1].text = k, v
    b2 = io.BytesIO(); d2.save(b2)
    inbox.files.update({"a/e7_SI.xlsx": xlsx("SHIPPING INSTRUCTION", xrows),
                        "a/e7_BL.docx": b2.getvalue(),
                        "a/e8_SI.xlsx": xlsx("Sheet1", xrows)})
    r7 = extract_email(inbox, {"email_id": "e7", "attachments": ["a/e7_SI.xlsx", "a/e7_BL.docx"]})
    r8 = extract_email(inbox, {"email_id": "e8", "attachments": ["a/e8_SI.xlsx", "a/e1_BL.txt"]})
    checks += [
        ("Excel title in sheet name -> SI", r7["si"] is not None),
        ("Word title in page header -> BL, despite 'Invoice No.' field", r7["bl"] is not None),
        ("'invoice' inside a field doesn't make it OTHER", not r7["other_documents"]),
        ("no title anywhere -> filename used, LLM asked", r8["si"] is not None and not is_confident(r8)),
        ("real invoice still OTHER", r2["other_documents"][0]["doc_type_detected"] == "OTHER"),
        ("SI or BL missing -> never confident", not is_confident(r2)),
    ]

    for name, ok in checks:
        print(f"{'PASS' if ok else 'FAIL'}  {name}")