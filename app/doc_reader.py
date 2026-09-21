"""
doc_reader.py - turns any attachment into plain text the extractors can read.

    .txt           -> read as-is
    .pdf (text)    -> pdfplumber: page text + tables
    .pdf (scanned) -> no text layer, so pages are rendered to images and the
                      LLM (vision) transcribes them
    .docx          -> python-docx: paragraphs + tables
    .xlsx          -> openpyxl: every non-empty row of every sheet
    .png / .jpg    -> LLM (vision) transcription

Table rows with exactly two cells are written as "Label: value", the same shape
as the plain-text SI/BL files, so the rule extractor can read them too.

Vision transcriptions are cached in output/doc_text/ so each scan is only sent
to the LLM once.

The easy way to use it is SmartInbox: it wraps the organizers' Inbox and makes
inbox.read_text() work for every format, so neither extractor needs changing:

    inbox = SmartInbox(Inbox(DATA_DIR), DATA_DIR)

Quick test on one file (from the CargoCheck folder):
    python app/doc_reader.py attachments/email_123_BL.pdf
"""
import json
import os
import re
import sys
from pathlib import Path

from dotenv import load_dotenv

load_dotenv()

BASE_DIR = Path(__file__).resolve().parent.parent
CACHE_DIR = BASE_DIR / "output" / "doc_text"
MIN_TEXT_CHARS = 40      # fewer letters/digits than this on a PDF = treat as a scan
MAX_VISION_PAGES = 5

VISION_PROMPT = """These images are pages of a scanned shipping document.
Transcribe ALL the text exactly as written.
- Keep each label with its value on one line, e.g. "Port of Loading: PORT KLANG (MYPKG)".
- For tables, write each row on one line with cells separated by " | ".
- Do not correct spelling, do not guess. Write [illegible] for parts you cannot read.
- Output only the transcription, no comments."""


# --- helpers -----------------------------------------------------------------

def _row_to_line(cells):
    """['Port of Loading', 'PORT KLANG'] -> 'Port of Loading: PORT KLANG'."""
    cells = [re.sub(r"\s+", " ", str(c)).strip() for c in cells if c is not None]
    cells = [c for c in cells if c]
    # merged cells often repeat the same value; drop consecutive duplicates
    deduped = [c for i, c in enumerate(cells) if i == 0 or c != cells[i - 1]]
    if not deduped:
        return ""
    if len(deduped) == 2 and not deduped[0].rstrip().endswith(":"):
        return f"{deduped[0]}: {deduped[1]}"
    if len(deduped) == 2:
        return f"{deduped[0]} {deduped[1]}"
    return " | ".join(deduped)


def _meaningful_chars(text):
    return len(re.sub(r"[^A-Za-z0-9]", "", text or ""))


# --- one reader per format ---------------------------------------------------

def read_pdf(path):
    import pdfplumber

    parts = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            parts.append(page.extract_text() or "")
            for table in page.extract_tables() or []:
                parts.extend(_row_to_line(row) for row in table)
    text = "\n".join(p for p in parts if p)

    if _meaningful_chars(text) >= MIN_TEXT_CHARS:
        return text, "pdf_text"

    # No usable text layer: it's a scan. Render pages and ask the LLM to read them.
    return transcribe_with_vision(path, _render_pdf_pages(path)), "pdf_vision"


def _render_pdf_pages(path):
    import io
    import pypdfium2 as pdfium

    images = []
    pdf = pdfium.PdfDocument(str(path))
    for i in range(min(len(pdf), MAX_VISION_PAGES)):
        pil = pdf[i].render(scale=2).to_pil()
        buf = io.BytesIO()
        pil.convert("RGB").save(buf, format="PNG")
        images.append(buf.getvalue())
    return images


def read_docx(path):
    import docx

    document = docx.Document(str(path))
    lines = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            lines.append(_row_to_line(cell.text for cell in row.cells))
    return "\n".join(l for l in lines if l), "docx"


def read_xlsx(path):
    import openpyxl

    workbook = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    lines = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(values_only=True):
            line = _row_to_line(row)
            if line:
                lines.append(line)
    workbook.close()
    return "\n".join(lines), "xlsx"


def read_image(path):
    return transcribe_with_vision(path, [Path(path).read_bytes()],
                                  fmt=Path(path).suffix.lower().lstrip(".")), "image_vision"


# --- vision (scans and images) -----------------------------------------------

def transcribe_with_vision(path, images, fmt="png"):
    """Send page images to the LLM and get the text back. Cached per file."""
    CACHE_DIR.mkdir(parents=True, exist_ok=True)
    cache = CACHE_DIR / (Path(path).name + ".txt")
    if cache.exists():
        return cache.read_text(encoding="utf-8")
    if not images:
        return ""

    fmt = "jpeg" if fmt == "jpg" else fmt
    provider = os.getenv("LLM_PROVIDER", "bedrock").lower()

    if provider == "bedrock":
        import boto3
        client = boto3.client("bedrock-runtime", region_name=os.getenv("AWS_REGION", "us-east-1"))
        content = [{"image": {"format": fmt, "source": {"bytes": img}}} for img in images]
        content.append({"text": VISION_PROMPT})
        response = client.converse(
            modelId=os.getenv("BEDROCK_MODEL_ID"),
            messages=[{"role": "user", "content": content}],
            inferenceConfig={"maxTokens": 4096, "temperature": 0},
        )
        blocks = response["output"]["message"]["content"]
        text = "".join(b["text"] for b in blocks if "text" in b)
    else:
        from google import genai
        from google.genai import types
        client = genai.Client()
        parts = [types.Part.from_bytes(data=img, mime_type=f"image/{fmt}") for img in images]
        response = client.models.generate_content(
            model=os.getenv("GEMINI_MODEL"),
            contents=parts + [VISION_PROMPT],
            config=types.GenerateContentConfig(temperature=0),
        )
        text = response.text or ""

    # Mostly illegible = genuinely unreadable; return empty so it goes to review.
    readable = re.sub(r"\[illegible\]", "", text, flags=re.I)
    if _meaningful_chars(readable) < MIN_TEXT_CHARS:
        text = ""

    cache.write_text(text, encoding="utf-8")
    return text


# --- main entry point ----------------------------------------------------------

READERS = {
    ".pdf": read_pdf,
    ".docx": read_docx,
    ".xlsx": read_xlsx,
    ".xlsm": read_xlsx,
    ".png": read_image,
    ".jpg": read_image,
    ".jpeg": read_image,
}


def read_document(path):
    """Return (text, method) for any supported file."""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".txt":
        return path.read_text(encoding="utf-8", errors="replace"), "txt"
    if suffix not in READERS:
        raise ValueError(f"Unsupported attachment type: {suffix}")
    return READERS[suffix](path)


class SmartInbox:
    """Wraps the organizers' Inbox so read_text() understands PDF, Word, Excel and scans.
    Everything else (iterating emails, submit, ...) is passed straight through."""

    def __init__(self, inbox, data_dir):
        self._inbox = inbox
        self._data_dir = Path(data_dir)

    def __iter__(self):
        return iter(self._inbox)

    def __getattr__(self, name):
        return getattr(self._inbox, name)

    def read_text(self, path):
        if Path(path).suffix.lower() == ".txt":
            return self._inbox.read_text(path)
        text, _method = read_document(self._data_dir / path)
        return text


if __name__ == "__main__":
    data_dir = Path(os.environ.get("DATA_DIR", str(BASE_DIR / "data" / "sdoc-hackathon-bundle")))
    for arg in sys.argv[1:]:
        text, method = read_document(data_dir / arg)
        print("=" * 70)
        print(f"{arg}  [{method}]  {len(text)} characters\n")
        print(text[:2000])